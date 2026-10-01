"""Reproduce the UCI retail portfolio analysis. No results are fabricated."""
from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import json
import sqlite3
import urllib.request
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCE_URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
SOURCE_PAGE = "https://archive.ics.uci.edu/dataset/352/online+retail"
ARCHIVE_SHA256 = "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"
WORKBOOK_SHA256 = "43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d"
HEADERS = ("InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country")


def identifier(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def pence(quantity, unit_price):
    # Round each extended line amount, not the unit price. Preserve subpenny prices.
    return int((Decimal(str(unit_price)) * quantity * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def download_source():
    raw = ROOT / "data" / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    archive = raw / "online-retail.zip"
    if not archive.exists():
        print("Downloading the original UCI dataset...", flush=True)
        request = urllib.request.Request(SOURCE_URL, headers={"User-Agent": "RetailPortfolio/1.0"})
        with urllib.request.urlopen(request, timeout=90) as response:
            archive.write_bytes(response.read())
    digest = hashlib.sha256(archive.read_bytes()).hexdigest()
    if digest != ARCHIVE_SHA256:
        raise ValueError("The source archive changed. Review its provenance before updating the pinned checksum.")
    with zipfile.ZipFile(archive) as source:
        (raw / "Online Retail.xlsx").write_bytes(source.read("Online Retail.xlsx"))
    return raw / "Online Retail.xlsx"


def fetch_rows(connection, filename):
    return [dict(row) for row in connection.execute((ROOT / "sql" / filename).read_text())]


def export_csv(path, rows):
    if not rows:
        return
    with path.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def analyze(input_path, output_dir, database_path):
    source_digest = hashlib.sha256(input_path.read_bytes()).hexdigest()
    if source_digest != WORKBOOK_SHA256:
        raise ValueError("The input workbook does not match the verified original UCI dataset.")
    output_dir.mkdir(parents=True, exist_ok=True)
    database_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(database_path)
    connection.row_factory = sqlite3.Row
    connection.execute("DROP TABLE IF EXISTS raw_lines")
    connection.execute("""CREATE TABLE raw_lines (
        source_row INTEGER PRIMARY KEY, invoice_no TEXT, stock_code TEXT,
        description TEXT, quantity INTEGER, invoice_date TEXT, unit_price TEXT,
        customer_id TEXT, country TEXT, line_value_pence INTEGER, date_valid INTEGER
    )""")
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    if tuple(next(iterator)) != HEADERS:
        raise ValueError("Source columns do not match the documented data dictionary.")
    batch = []
    seen = set()
    independent = Counter()
    independent_products = defaultdict(int)
    raw_quality = Counter()
    insert_sql = "INSERT INTO raw_lines VALUES (?,?,?,?,?,?,?,?,?,?,?)"
    for source_row, row in enumerate(iterator, start=2):
        invoice, stock, description, quantity, date, price, customer, country = row
        if quantity is None or price is None:
            raise ValueError(f"Missing quantity or price on source row {source_row}.")
        quantity = int(quantity)
        valid_date = isinstance(date, datetime)
        date_text = date.isoformat(sep=" ") if valid_date else str(date)
        original_key = (identifier(invoice), identifier(stock), description, quantity,
                        date_text, str(price), identifier(customer), country)
        amount = pence(quantity, price)
        batch.append((source_row, *original_key, amount, int(valid_date)))
        raw_quality["rows"] += 1
        raw_quality["missing_customer"] += customer is None
        raw_quality["missing_description"] += description is None
        raw_quality["cancel_flag"] += str(invoice).upper().startswith("C")
        # Independent Python control totals. SQL performs its own deduplication.
        duplicate = original_key in seen
        seen.add(original_key)
        valid = (not duplicate and valid_date and Decimal(str(price)) > 0 and quantity != 0
                 and not (str(invoice).upper().startswith("C") and quantity > 0))
        if valid:
            independent["accepted_lines"] += 1
            independent["net_pence"] += amount
            independent["gross_pence" if quantity > 0 else "negative_pence"] += amount
            if str(stock)[:1].isdigit():
                independent_products[identifier(stock)] += amount
        if len(batch) == 5000:
            connection.executemany(insert_sql, batch)
            batch.clear()
        if raw_quality["rows"] % 100000 == 0:
            print(f"Read {raw_quality['rows']:,} invoice lines...", flush=True)
    if batch:
        connection.executemany(insert_sql, batch)
    workbook.close()
    connection.commit()
    connection.executescript((ROOT / "sql" / "01_model.sql").read_text())
    kpis = fetch_rows(connection, "02_kpis.sql")[0]
    monthly = fetch_rows(connection, "03_monthly.sql")
    products = fetch_rows(connection, "04_products.sql")
    countries = fetch_rows(connection, "05_countries.sql")
    dispositions = fetch_rows(connection, "06_quality.sql")

    start_date = datetime.fromisoformat(kpis["start_at"])
    end_date = datetime.fromisoformat(kpis["end_at"])
    for row in monthly:
        year, month = map(int, row["month"].split("-"))
        last_day = calendar.monthrange(year, month)[1]
        row["partial"] = ((row["month"] == start_date.strftime("%Y-%m") and start_date.day > 1)
                          or (row["month"] == end_date.strftime("%Y-%m") and end_date.day < last_day))
        if row["partial"]:
            row["gross_month_change_pct"] = None
    # All country/month combinations let the static dashboard filter exact aggregates.
    country_month = [dict(row) for row in connection.execute("""
        SELECT country, month,
               SUM(CASE WHEN line_type = 'sale' THEN line_value_pence ELSE 0 END) AS gross_pence,
               -SUM(CASE WHEN line_type = 'credit' THEN line_value_pence ELSE 0 END) AS credit_pence,
               SUM(line_value_pence) AS net_pence,
               COUNT(DISTINCT CASE WHEN line_type = 'sale' THEN invoice_no END) AS sale_invoices
        FROM invoice_lines GROUP BY country, month ORDER BY country, month
    """)]
    # Product/month/country aggregates support product rankings within the same filter.
    # To keep the website light, only the overall leading 20 merchandise-like SKUs
    # are exported. The complete SKU rankings remain in the downloadable CSV.
    top_codes = [row["stock_code"] for row in products[:20]]
    product_segments = [dict(row) for row in connection.execute("""
        SELECT stock_code, country, month,
               SUM(CASE WHEN line_type = 'sale' THEN line_value_pence ELSE 0 END) AS gross_pence,
               -SUM(CASE WHEN line_type = 'credit' THEN line_value_pence ELSE 0 END) AS credit_pence,
               SUM(line_value_pence) AS net_pence
        FROM invoice_lines WHERE stock_code IN (%s)
        GROUP BY stock_code, country, month ORDER BY stock_code, country, month
    """ % ",".join("?" for _ in top_codes), top_codes)]

    checks = {}
    def check(name, condition):
        checks[name] = bool(condition)
        if not condition:
            raise AssertionError(f"Validation failed: {name}")
    check("source_row_count", raw_quality["rows"] == 541909)
    check("row_disposition_reconciliation", sum(row["rows"] for row in dispositions) == raw_quality["rows"])
    for key in ("accepted_lines", "gross_pence", "net_pence"):
        check("independent_python_" + key, kpis[key] == independent[key])
    check("independent_python_credits", kpis["credit_pence"] == -independent["negative_pence"])
    check("gross_less_credits_equals_net", kpis["gross_pence"] - kpis["credit_pence"] == kpis["net_pence"])
    check("monthly_net_reconciliation", sum(row["net_pence"] for row in monthly) == kpis["net_pence"])
    check("country_net_reconciliation", sum(row["net_pence"] for row in countries) == kpis["net_pence"])
    check("country_month_reconciliation", sum(row["net_pence"] for row in country_month) == kpis["net_pence"])
    check("country_month_invoices_no_double_count", sum(row["sale_invoices"] for row in country_month) == kpis["sale_invoices"])
    check("complete_product_control", all(row["net_pence"] == independent_products[row["stock_code"]] for row in products))
    check("product_join_no_fanout", len(products) == len(independent_products))
    check("partial_final_month_marked", monthly[-1]["partial"] and monthly[-1]["gross_month_change_pct"] is None)
    duplicate_value = sum(row["signed_pence"] or 0 for row in dispositions if row["disposition"] == "duplicate")
    sensitivity = dict(connection.execute("""
        SELECT SUM(line_value_pence) AS raw_valid_net_pence FROM raw_lines
        WHERE date_valid = 1 AND CAST(unit_price AS REAL) > 0 AND quantity != 0
        AND NOT (UPPER(invoice_no) LIKE 'C%' AND quantity > 0)
    """).fetchone())
    result = {
        "source": {"title": "UCI Online Retail", "url": SOURCE_PAGE,
                   "doi": "https://doi.org/10.24432/C5BW33", "license": "CC BY 4.0",
                   "attribution": "Chen, D. (2015). Online Retail. UCI Machine Learning Repository.",
                   "archive_sha256": ARCHIVE_SHA256,
                   "workbook_sha256": source_digest},
        "units": "GBP; integer pence, rounded half up per invoice line",
        "raw_quality": dict(raw_quality), "kpis": kpis, "monthly": monthly,
        "countries": countries, "top_products": products[:20],
        "country_month": country_month, "product_segments": product_segments,
        "dispositions": dispositions, "sensitivity": sensitivity,
        "validation": {"checks": checks, "passed": all(checks.values()), "check_count": len(checks)},
        "implementation": "AI-assisted portfolio study prepared for Justin Pierite. Historical public data; no claimed employer or commercial impact."
    }
    (output_dir / "analysis.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    for name, rows in [("monthly", monthly), ("products", products), ("countries", countries), ("quality", dispositions)]:
        export_csv(output_dir / f"{name}.csv", rows)
    (output_dir / "validation.json").write_text(json.dumps(result["validation"], indent=2), encoding="utf-8")
    connection.close()
    print(json.dumps({"raw_rows": raw_quality["rows"], "kpis": kpis, "dispositions": dispositions,
                      "top_product": products[0], "top_country": countries[0],
                      "validation": result["validation"], "sensitivity": sensitivity}, indent=2), flush=True)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, help="Existing original Online Retail.xlsx; otherwise downloads the pinned UCI source.")
    args = parser.parse_args()
    analyze(args.input or download_source(), ROOT / "results", ROOT / "data" / "derived" / "retail.sqlite")
